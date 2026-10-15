"""
Word出力モジュール
バイアス評価表 / PRISMAフロー図 / メタアナリシスレポートを.docx形式で生成
"""

import io
from datetime import date
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ===== 共通ユーティリティ =====

def _set_cell_bg(cell, hex_color: str):
    """セル背景色を設定する"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _set_cell_font(cell, text: str, bold: bool = False, size_pt: int = 10,
                   color_hex: str = None):
    """セルのテキストとフォントを設定する"""
    cell.text = text
    run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.paragraphs[0].add_run(text)
    run.text = text
    run.bold = bold
    run.font.size = Pt(size_pt)
    if color_hex:
        r, g, b = int(color_hex[0:2], 16), int(color_hex[2:4], 16), int(color_hex[4:6], 16)
        run.font.color.rgb = RGBColor(r, g, b)


# ===== バイアス評価表 =====

RATING_BG = {
    "Low":      "D1FAE5",  # 薄緑
    "Moderate": "FEF3C7",  # 薄黄
    "High":     "FEE2E2",  # 薄赤
    "Unclear":  "F3F4F6",  # 薄グレー
}

RATING_TEXT_COLOR = {
    "Low":      "065F46",
    "Moderate": "92400E",
    "High":     "991B1B",
    "Unclear":  "374151",
}

DOMAIN_LABELS = [
    "Study Participation",
    "Study Attrition",
    "Prognostic Factor Measurement",
    "Outcome Measurement",
    "Study Confounding",
    "Statistical Analysis",
]


def export_bias_table(bias_results: list) -> io.BytesIO:
    """
    バイアス評価表をWordドキュメントとして出力する

    Parameters
    ----------
    bias_results : list of BiasResult
    """
    doc = Document()

    # ページ余白
    section = doc.sections[0]
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    # タイトル
    title = doc.add_heading("Risk of Bias Assessment (QUIPS)", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Generated: {date.today().strftime('%Y-%m-%d')}")
    doc.add_paragraph()

    if not bias_results:
        doc.add_paragraph("バイアス評価データがありません。")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        return buf

    # 表作成: 行=論文、列=ドメイン+総合
    n_cols = len(DOMAIN_LABELS) + 2  # 論文名 + 6ドメイン + 総合
    table = doc.add_table(rows=1 + len(bias_results), cols=n_cols)
    table.style = "Table Grid"

    # ヘッダー行
    hdr = table.rows[0].cells
    hdr[0].text = "Study"
    hdr[0].paragraphs[0].runs[0].bold = True
    for j, label in enumerate(DOMAIN_LABELS):
        hdr[j + 1].text = label
        run = hdr[j + 1].paragraphs[0].runs[0]
        run.bold = True
        run.font.size = Pt(8)
    hdr[-1].text = "Overall"
    hdr[-1].paragraphs[0].runs[0].bold = True

    # ヘッダー背景
    for cell in hdr:
        _set_cell_bg(cell, "1E3A5F")
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # データ行
    for i, bias_res in enumerate(bias_results):
        row = table.rows[i + 1].cells

        # 論文名
        short_title = bias_res.title[:40] + ("..." if len(bias_res.title) > 40 else "")
        row[0].text = f"{short_title}\n(PMID: {bias_res.pmid})"
        row[0].paragraphs[0].runs[0].font.size = Pt(8)

        # ドメイン評価
        for j, domain in enumerate(bias_res.domains):
            from bias_assessment import get_confirmed_rating
            rating = get_confirmed_rating(domain)
            cell = row[j + 1]
            cell.text = rating
            _set_cell_bg(cell, RATING_BG.get(rating, "FFFFFF"))
            run = cell.paragraphs[0].runs[0]
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(
                *[int(RATING_TEXT_COLOR.get(rating, "000000")[k:k+2], 16) for k in (0, 2, 4)]
            )
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # 総合評価
        overall = bias_res.overall
        _set_cell_bg(row[-1], RATING_BG.get(overall, "FFFFFF"))
        row[-1].text = overall
        run = row[-1].paragraphs[0].runs[0]
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(
            *[int(RATING_TEXT_COLOR.get(overall, "000000")[k:k+2], 16) for k in (0, 2, 4)]
        )
        row[-1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row[-1].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # 列幅調整
    for row in table.rows:
        row.cells[0].width = Cm(5)
        for j in range(1, n_cols - 1):
            row.cells[j].width = Cm(3.2)
        row.cells[-1].width = Cm(2.5)

    doc.add_paragraph()
    doc.add_paragraph("Note: Ratings based on QUIPS tool (Quality In Prognosis Studies).")
    doc.add_paragraph(
        "Domains: Study Participation, Study Attrition, Prognostic Factor Measurement, "
        "Outcome Measurement, Study Confounding, Statistical Analysis."
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ===== PRISMAフロー図 =====

def export_prisma_flow(
    n_identified: int,
    n_duplicates: int,
    n_screened: int,
    n_excluded_title: int,
    n_fulltext: int,
    n_excluded_fulltext: int,
    n_included: int,
    pico: dict = None,
) -> io.BytesIO:
    """
    PRISMA 2020フロー図をWordドキュメントとして出力する
    （ASCII形式のテーブルレイアウト）
    """
    doc = Document()

    title = doc.add_heading("PRISMA 2020 Flow Diagram", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if pico:
        doc.add_paragraph(
            f"Population: {pico.get('population', '')}  |  "
            f"Sport: {pico.get('sport', '')}  |  "
            f"Injury: {pico.get('injury', '')}"
        ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph(f"Generated: {date.today().strftime('%Y-%m-%d')}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # --- PRISMA フロー表 ---
    table = doc.add_table(rows=12, cols=3)
    table.style = "Table Grid"

    def fill_box(row_idx, col_idx, text, bg="DBEAFE", bold=False):
        cell = table.cell(row_idx, col_idx)
        cell.text = text
        _set_cell_bg(cell, bg)
        for para in cell.paragraphs:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.bold = bold
                run.font.size = Pt(9)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    def arrow(row_idx, col_idx):
        cell = table.cell(row_idx, col_idx)
        cell.text = "↓"
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_cell_bg(cell, "FFFFFF")

    def exclusion_box(row_idx, col_idx, text):
        cell = table.cell(row_idx, col_idx)
        cell.text = text
        _set_cell_bg(cell, "FEF3C7")
        for para in cell.paragraphs:
            for run in para.runs:
                run.font.size = Pt(9)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Identification
    fill_box(0, 1, f"Records identified\nPubMed (n = {n_identified})", bg="DBEAFE", bold=True)
    arrow(1, 1)

    # After duplicates
    fill_box(2, 1, f"Records after duplicates removed\n(n = {n_screened})\n\nDuplicates removed: {n_duplicates}", bg="DBEAFE")
    exclusion_box(2, 2, f"Duplicates removed\n(n = {n_duplicates})")
    arrow(3, 1)

    # Screening
    fill_box(4, 1, f"Records screened\n(n = {n_screened})", bg="EDE9FE")
    exclusion_box(4, 2, f"Records excluded\n(n = {n_excluded_title})\nNot meeting inclusion criteria")
    arrow(5, 1)

    # Full-text
    fill_box(6, 1, f"Full-text articles assessed\n(n = {n_fulltext})", bg="D1FAE5")
    exclusion_box(6, 2, f"Full-text excluded\n(n = {n_excluded_fulltext})\nFull text not available / not relevant")
    arrow(7, 1)

    # Included
    fill_box(8, 1, f"Studies included\nin synthesis (n = {n_included})", bg="22C55E", bold=True)
    for col in range(3):
        table.cell(9, col).text = ""
        table.cell(10, col).text = ""
        table.cell(11, col).text = ""

    # 列幅
    for row in table.rows:
        row.cells[0].width = Cm(2)
        row.cells[1].width = Cm(10)
        row.cells[2].width = Cm(8)
        row.height = Cm(1.5)

    doc.add_paragraph()
    doc.add_paragraph(
        "Note: Flow diagram based on PRISMA 2020 guidelines "
        "(Page MJ, et al. BMJ 2021;372:n71)."
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ===== メタアナリシスレポート =====

def export_meta_report(
    meta_result,
    studies: list,
    selected_factor: str,
    forest_plot_buf: "io.BytesIO",
    pico: dict = None,
) -> io.BytesIO:
    """
    メタアナリシス結果レポートをWordドキュメントとして出力する
    """
    from meta_analysis import interpret_heterogeneity, heterogeneity_recommendation

    doc = Document()

    # タイトル
    title = doc.add_heading("Meta-Analysis Report", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if pico:
        doc.add_paragraph(
            f"Sport: {pico.get('sport', '')}  |  Injury: {pico.get('injury', '')}  |  "
            f"Risk Factor: {selected_factor}"
        ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"Generated: {date.today().strftime('%Y-%m-%d')}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # 1. 方法
    doc.add_heading("1. Methods", level=2)
    method_full = {
        "MH": "Mantel-Haenszel",
        "IV": "Inverse Variance",
        "Peto": "Peto",
    }.get(meta_result.method, meta_result.method)
    var_full = {
        "REML": "Restricted Maximum Likelihood (REML)",
        "DL": "DerSimonian-Laird (DL)",
        "fixed": "Fixed Effects",
    }.get(meta_result.variance_method, meta_result.variance_method)

    doc.add_paragraph(
        f"A random-effects meta-analysis was conducted using the {method_full} method "
        f"with {var_full} heterogeneity estimation. "
        f"The outcome measure was {meta_result.measure_type}. "
        f"Statistical analyses were performed using a custom Python implementation."
    )

    # 2. 結果
    doc.add_heading("2. Results", level=2)

    # 結果サマリー表
    result_table = doc.add_table(rows=2, cols=6)
    result_table.style = "Table Grid"

    headers = ["n Studies", "Pooled Effect", "95% CI", "I²", "τ²", "Q (p-value)"]
    for j, h in enumerate(headers):
        cell = result_table.rows[0].cells[j]
        cell.text = h
        _set_cell_bg(cell, "1E3A5F")
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    vals = [
        str(meta_result.n_studies),
        f"{meta_result.pooled_effect:.2f}",
        f"{meta_result.ci_lower:.2f} – {meta_result.ci_upper:.2f}",
        f"{meta_result.i2:.1f}%",
        f"{meta_result.tau2:.4f}",
        f"{meta_result.q_stat:.2f} (p={meta_result.q_pvalue:.3f})",
    ]
    for j, v in enumerate(vals):
        result_table.rows[1].cells[j].text = v
        result_table.rows[1].cells[j].paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph()

    # 異質性の解釈
    h_label, h_icon = interpret_heterogeneity(meta_result.i2)
    doc.add_paragraph(
        f"Heterogeneity was {h_label.lower()} (I² = {meta_result.i2:.1f}%). "
        f"{heterogeneity_recommendation(meta_result.i2)}"
    )

    # 3. 個別研究
    doc.add_heading("3. Individual Studies", level=2)
    study_table = doc.add_table(rows=1 + len(studies), cols=5)
    study_table.style = "Table Grid"

    s_headers = ["Study", "Effect", "CI Lower", "CI Upper", "Weight (%)"]
    for j, h in enumerate(s_headers):
        cell = study_table.rows[0].cells[j]
        cell.text = h
        _set_cell_bg(cell, "1E3A5F")
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(255, 255, 255)

    for i, s in enumerate(studies):
        row = study_table.rows[i + 1].cells
        row[0].text = s.study_id
        row[1].text = f"{s.effect:.2f}"
        row[2].text = f"{s.ci_lower:.2f}" if s.ci_lower else "-"
        row[3].text = f"{s.ci_upper:.2f}" if s.ci_upper else "-"
        row[4].text = f"{s.weight:.1f}"
        for cell in row:
            cell.paragraphs[0].runs[0].font.size = Pt(9)

    # 4. フォレストプロット
    doc.add_heading("4. Forest Plot", level=2)
    if forest_plot_buf:
        forest_plot_buf.seek(0)
        doc.add_picture(forest_plot_buf, width=Inches(6.5))
        last_para = doc.paragraphs[-1]
        last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()
    doc.add_paragraph(
        f"Figure. Forest plot of {meta_result.measure_type} for the association between "
        f"'{selected_factor}' and injury risk. "
        f"Boxes represent individual study estimates (sized by weight); "
        f"diamond represents the pooled estimate."
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
